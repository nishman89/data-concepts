"""Generates the Sparta Leisure dataset used by the whole site.

Run once (python3 scripts/generate_data.py) and commit what it writes. The fixed seed means the same files every time,
so every exercise answer is stable. It writes:

  data/workbook.json                         the lesson data, embedded in the site (dates are Excel serial numbers)
  data/capstone_facts.json                   numbers the capstone page quotes ("check your numbers")
  public/downloads/sparta-leisure.xlsx   the same lesson data as an Excel workbook
  public/downloads/sparta-leisure-capstone.zip   the raw, messy capstone export (CSV, JSON and XML)

Needs numpy and openpyxl.
"""
import csv, io, json, math, os, random, zipfile
from datetime import date, datetime, timedelta
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
rng = np.random.default_rng(2026)
random.seed(2026)

EPOCH = date(1899, 12, 30)
def serial(d):
    if isinstance(d, datetime):
        return (d - datetime(1899, 12, 30)).total_seconds() / 86400
    return (d - EPOCH).days

# ---------------------------------------------------------------- reference tables
SITES = [  # id, name, town, pool, capacity, opened, lat, lon, demand weight
    ('S01', 'Telford Central', 'Telford', True, 220, date(2009, 4, 1), 52.6766, -2.4469, 1.25),
    ('S02', 'Shrewsbury Riverside', 'Shrewsbury', True, 180, date(2012, 9, 3), 52.7073, -2.7553, 1.0),
    ('S03', 'Wolverhampton City', 'Wolverhampton', False, 250, date(2015, 1, 12), 52.5862, -2.1288, 1.35),
    ('S04', 'Walsall Arboretum', 'Walsall', False, 140, date(2018, 6, 4), 52.5893, -1.9748, 0.8),
    ('S05', 'Dudley Castle Gate', 'Dudley', True, 200, date(2016, 3, 14), 52.5120, -2.0811, 1.0),
    ('S06', 'Kidderminster Wyre', 'Kidderminster', False, 120, date(2021, 10, 4), 52.3885, -2.2497, 0.6),
]
SITE_IDS = [s[0] for s in SITES]
POOL = {s[0]: s[3] for s in SITES}
POSTCODES = {'S01': ['TF1', 'TF2', 'TF3', 'TF4', 'TF7'], 'S02': ['SY1', 'SY2', 'SY3', 'SY4'],
             'S03': ['WV1', 'WV2', 'WV3', 'WV4', 'WV6'], 'S04': ['WS1', 'WS2', 'WS3', 'WS5'],
             'S05': ['DY1', 'DY2', 'DY3', 'DY5'], 'S06': ['DY10', 'DY11', 'DY14']}

PLANS = [  # code, name, monthly fee, joining fee, peak access, classes per month, pool access, min age
    ('OFF', 'Off-peak', 21.99, 10, False, 4, False, 16),
    ('STD', 'Standard', 32.99, 15, True, 8, True, 16),
    ('PREM', 'Premium', 44.99, 0, True, 20, True, 16),
    ('STU', 'Student', 18.99, 0, False, 4, True, 16),
    ('SEN', 'Senior', 19.99, 0, False, 8, True, 60),
]
PLAN_FEE = {p[0]: p[2] for p in PLANS}
PRICE_YEARS = [2022, 2023, 2024, 2025]
PRICE_HISTORY = {'OFF': [18.99, 19.99, 20.99, 21.99], 'STD': [28.99, 29.99, 31.49, 32.99],
                 'PREM': [39.99, 41.99, 42.99, 44.99], 'STU': [15.99, 16.99, 17.99, 18.99], 'SEN': [16.99, 17.99, 18.99, 19.99]}

CLASSES = [  # id, name, category, duration, intensity, max places, needs pool
    ('C01', 'Spin', 'Cardio', 45, 'High', 20, False), ('C02', 'HIIT', 'Cardio', 30, 'High', 18, False),
    ('C03', 'Body Pump', 'Strength', 60, 'Medium', 24, False), ('C04', 'Kettlebells', 'Strength', 45, 'High', 14, False),
    ('C05', 'Yoga', 'Mind & Body', 60, 'Low', 20, False), ('C06', 'Pilates', 'Mind & Body', 45, 'Low', 16, False),
    ('C07', 'Zumba', 'Cardio', 45, 'Medium', 25, False), ('C08', 'Circuits', 'Strength', 45, 'Medium', 20, False),
    ('C09', 'Aqua Fit', 'Aqua', 45, 'Low', 22, True), ('C10', 'Lane Swim Coaching', 'Aqua', 60, 'Medium', 10, True),
    ('C11', 'Boxercise', 'Cardio', 45, 'High', 16, False), ('C12', 'Stretch & Mobility', 'Mind & Body', 30, 'Low', 20, False),
]
CLASS_TIMES = ['06:30', '07:15', '09:30', '12:15', '17:30', '18:30', '19:30']

FIRST_F = 'Olivia Amelia Isla Ava Mia Sophia Grace Lily Freya Emily Ella Poppy Evie Chloe Ruby Alice Hannah Jessica Priya Aisha Zara Megan Holly Sarah Laura Rachel Emma Charlotte Lucy Katie Jade Amira Nadia Fatima Gemma Helen Claire Joanne Susan Margaret'.split()
FIRST_M = 'Oliver George Harry Noah Jack Leo Arthur Oscar Charlie Jacob Thomas Henry William James Mohammed Alfie Joshua Ethan Daniel Samuel Ryan Liam Adam Ben Callum Kieran Rhys Aaron Imran Hassan Ravi Arjun Tom Matthew Mark Paul Stephen Gary Andrew David'.split()
LAST = 'Smith Jones Taylor Brown Williams Wilson Johnson Davies Patel Robinson Wright Thompson Evans Walker White Roberts Green Hall Wood Jackson Clarke Khan Hughes Edwards Turner Hill Moore Cooper Ward Morris King Baker Harris Lewis Price Begum Singh Kaur Bennett Griffiths Lloyd Morgan Powell Shaw Fisher Mason Webb Ali Hussain'.split()

MONTH_TEMP = [4.6, 5.2, 7.4, 9.9, 13.1, 16.0, 18.1, 17.7, 15.0, 11.4, 7.6, 5.0]
MONTH_DEMAND = [1.32, 1.12, 1.05, 0.98, 0.96, 0.9, 0.86, 0.82, 0.98, 1.0, 0.95, 0.78]   # visits by month
JOIN_SEASON = [1.9, 1.2, 1.1, 0.9, 0.85, 0.75, 0.7, 0.72, 1.2, 1.0, 0.85, 0.75]       # joiners by month


def month_start(y, m): return date(y, m, 1)
def month_end(y, m): return (date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1))
def rand_day(y, m): return date(y, m, int(rng.integers(1, month_end(y, m).day + 1)))

# ---------------------------------------------------------------- members
members = []
def new_member(site, join):
    gender = rng.choice(['F', 'M', 'X'], p=[0.52, 0.46, 0.02])
    age_at_join = float(np.clip(rng.normal(37, 13), 16, 82))
    plan = rng.choice(['OFF', 'STD', 'PREM', 'STU', 'SEN'], p=[0.22, 0.4, 0.2, 0.12, 0.06])
    if age_at_join >= 60 and rng.random() < 0.55: plan = 'SEN'
    if plan == 'SEN' and age_at_join < 60: age_at_join = float(rng.uniform(60, 78))
    if plan == 'STU': age_at_join = float(np.clip(rng.normal(20.5, 2.2), 16, 30))
    dob = join - timedelta(days=int(age_at_join * 365.25))
    first = rng.choice(FIRST_F if gender == 'F' else FIRST_M if gender == 'M' else FIRST_F + FIRST_M)
    m = dict(id=None, first=str(first), last=str(rng.choice(LAST)), gender=str(gender), dob=dob,
             postcode=str(rng.choice(POSTCODES[site])), site=site, plan=str(plan), join=join, cancel=None,
             keen=float(rng.gamma(2.2, 0.5)))  # visits a week, before seasonality
    # Cancelling: students and off-peak churn most; the January joiners churn faster
    hazard = {'OFF': 0.026, 'STD': 0.018, 'PREM': 0.012, 'STU': 0.04, 'SEN': 0.01}[m['plan']] * (1.5 if join.month == 1 else 1)
    life = rng.exponential(1 / hazard) * 30.4
    c = join + timedelta(days=int(life) + 20)
    if c <= date(2026, 9, 30): m['cancel'] = c
    members.append(m)
    return m

# Joiners from 2019 to 2024: background population, weighted by site and season
for y in range(2019, 2025):
    for mo in range(1, 13):
        for s in SITES:
            if month_end(y, mo) < s[5]: continue
            for _ in range(rng.poisson(5.2 * s[8] * JOIN_SEASON[mo - 1])):
                new_member(s[0], rand_day(y, mo))

# 2025: joiners depend on marketing spend (so the correlation in SiteMonths is real)
spend = {}
for s in SITES:
    for mo in range(1, 13):
        base = 420 * s[8] * (1.9 if mo == 1 else 1.35 if mo == 9 else 1.0)
        spend[(s[0], mo)] = round(float(base * rng.uniform(0.7, 1.3)) / 10) * 10
        lam = (spend[(s[0], mo)] / 100) * 1.25 * JOIN_SEASON[mo - 1] ** 0.6 + rng.normal(0, 1)
        for _ in range(max(0, int(rng.poisson(max(lam, 0.5))))):
            new_member(s[0], rand_day(2025, mo))
# 2026 (January to September): for the capstone only
spend26 = {}
for s in SITES:
    for mo in range(1, 10):
        base = 440 * s[8] * (2.6 if mo == 1 else 1.35 if mo == 9 else 1.0)  # the big January campaign
        spend26[(s[0], mo)] = round(float(base * rng.uniform(0.75, 1.25)) / 10) * 10
        lam = (spend26[(s[0], mo)] / 100) * 1.1 * JOIN_SEASON[mo - 1] ** 0.75
        for _ in range(int(rng.poisson(lam))):
            new_member(s[0], rand_day(2026, mo))

members.sort(key=lambda m: (m['join'], m['last']))
for i, m in enumerate(members, 1): m['id'] = f'M{i:04d}'

END25 = date(2025, 12, 31)
def status_at(m, d):
    return 'Cancelled' if (m['cancel'] and m['cancel'] <= d) else 'Active'

lesson_members = [m for m in members if m['join'] <= END25]

# ---------------------------------------------------------------- visits and bookings
def active_on(m, d): return m['join'] <= d and (m['cancel'] is None or d < m['cancel'])
def checkin_time(m, d):
    weekend = d.weekday() >= 5
    if PLANS[[p[0] for p in PLANS].index(m['plan'])][4] or weekend:
        hour = float(rng.choice([6.5, 7.5, 9, 10.5, 12.5, 14, 16.5, 17.5, 18.5, 19.5, 20.5],
                                p=[.07, .09, .07, .07, .09, .05, .08, .15, .16, .11, .06]))
    else:  # off-peak plans can only come in before 16:00 on weekdays
        hour = float(rng.choice([6.5, 7.5, 9, 10.5, 12.5, 14, 15], p=[.1, .1, .2, .2, .2, .12, .08]))
    mins = int(hour * 60 + rng.integers(0, 50))
    return datetime(d.year, d.month, d.day, mins // 60, mins % 60)

def make_visits(start, end, scale):
    out = []
    days = (end - start).days + 1
    for m in members:
        a = max(start, m['join']); b = min(end, (m['cancel'] - timedelta(days=1)) if m['cancel'] else end)
        if b < a: continue
        weeks = ((b - a).days + 1) / 7
        n = rng.poisson(m['keen'] * weeks * scale)
        for _ in range(n):
            d = a + timedelta(days=int(rng.integers(0, (b - a).days + 1)))
            if rng.random() > MONTH_DEMAND[d.month - 1] / 1.32: continue  # seasonality by thinning
            site = m['site'] if rng.random() < 0.9 else str(rng.choice(SITE_IDS))
            dur = int(round(float(np.clip(rng.lognormal(math.log(62), 0.42), 12, 240))))
            out.append(dict(member=m['id'], site=site, checkin=checkin_time(m, d), dur=dur, m=m))
    out.sort(key=lambda v: v['checkin'])
    return out

visits25 = make_visits(date(2025, 1, 1), END25, 0.255)
# A handful of people forgot to check out: real outliers to find in the statistics unit
for v in rng.choice(len(visits25), 9, replace=False):
    visits25[int(v)]['dur'] = int(rng.integers(420, 720))
for i, v in enumerate(visits25, 1): v['id'] = f'V{i:05d}'

def make_bookings(start, end, scale):
    out = []
    for m in members:
        a = max(start, m['join']); b = min(end, (m['cancel'] - timedelta(days=1)) if m['cancel'] else end)
        if b < a: continue
        fan = rng.random() < 0.55
        if not fan: continue
        favs = rng.choice(len(CLASSES), 2, replace=False)
        n = rng.poisson(scale * {'PREM': 1.8, 'STD': 1.1, 'SEN': 1.2, 'OFF': 0.7, 'STU': 0.8}[m['plan']] * ((b - a).days + 1) / 30.4)
        for _ in range(n):
            c = CLASSES[int(favs[0] if rng.random() < 0.6 else favs[1] if rng.random() < 0.6 else rng.integers(0, len(CLASSES)))]
            site = m['site']
            if c[6] and not POOL[site]:
                site = 'S01' if site in ('S03', 'S04') else 'S05' if site == 'S06' else 'S02'
            d = a + timedelta(days=int(rng.integers(0, (b - a).days + 1)))
            att = rng.random() < (0.9 if m['plan'] in ('PREM', 'SEN') else 0.78 if m['plan'] == 'STU' else 0.85)
            out.append(dict(member=m['id'], cls=c[0], site=site, date=d, time=str(rng.choice(CLASS_TIMES)), att=bool(att)))
    out.sort(key=lambda x: (x['date'], x['time']))
    return out

bookings25 = make_bookings(date(2025, 1, 1), END25, 0.62)
for i, b in enumerate(bookings25, 1): b['id'] = f'B{i:05d}'

# ---------------------------------------------------------------- SiteMonths (2025)
sitemonths = []
for s in SITES:
    for mo in range(1, 13):
        joins = sum(1 for m in members if m['site'] == s[0] and m['join'].year == 2025 and m['join'].month == mo)
        cancels = sum(1 for m in members if m['site'] == s[0] and m['cancel'] and m['cancel'].year == 2025 and m['cancel'].month == mo)
        vis = sum(1 for v in visits25 if v['site'] == s[0] and v['checkin'].month == mo)
        temp = round(MONTH_TEMP[mo - 1] + float(rng.normal(0, 0.8)), 1)
        sitemonths.append([s[0], month_start(2025, mo), joins, cancels, vis, spend[(s[0], mo)], temp])

# ---------------------------------------------------------------- MembersRaw: a small messy extract for the cleaning unit
raw_src = [m for m in lesson_members if m['site'] in ('S01', 'S02')][:48]
def messy_name(s, i):
    return [s, s.upper(), s.lower(), '  ' + s, s + ' ', s][i % 6]
raw_rows = []
for i, m in enumerate(raw_src):
    dob = m['dob']
    dob_txt = [dob.strftime('%d/%m/%Y'), dob.isoformat(), dob.strftime('%d %b %Y'), dob.strftime('%d/%m/%Y')][i % 4]
    gender = [m['gender'], {'F': 'Female', 'M': 'Male', 'X': 'X'}[m['gender']], m['gender'].lower(), m['gender']][i % 4]
    phone = '07' + ''.join(str(int(x)) for x in rng.integers(0, 10, 9))
    row = [m['id'], messy_name(m['first'], i), messy_name(m['last'], i + 3), gender, dob_txt, int(phone), m['postcode'], m['site'], m['plan'], round(PLAN_FEE[m['plan']], 2)]
    raw_rows.append(row)
# planted problems
raw_rows[5][3] = ''; raw_rows[17][3] = ''; raw_rows[30][6] = ''; raw_rows[41][6] = ''; raw_rows[22][9] = None
raw_rows[9][4] = '01/01/1900'; raw_rows[26][9] = 3299  # fee typed in pence
raw_rows[12][8] = 'Standard'; raw_rows[35][8] = 'std'
raw_rows[19][7] = 'Telford'
raw_rows.insert(14, list(raw_rows[13])); raw_rows.insert(33, list(raw_rows[31])); raw_rows.insert(40, list(raw_rows[2]))

# ---------------------------------------------------------------- workbook model (shared by the site and the xlsx)
def col(name, t='text', w=12, fmt=None): return dict(name=name, type=t, width=w, fmt=fmt)
SHEETS = []
SHEETS.append(dict(name='Sites', note='One row per leisure centre.', cols=[col('SiteID', w=8), col('SiteName', w=22), col('Town', w=15), col('HasPool', 'bool', 9), col('Capacity', 'int', 9), col('Opened', 'date', 11), col('Latitude', 'num', 10, '0.0000'), col('Longitude', 'num', 10, '0.0000')],
                   rows=[[s[0], s[1], s[2], s[3], s[4], serial(s[5]), s[6], s[7]] for s in SITES]))
SHEETS.append(dict(name='Plans', note='One row per membership plan. Fees are in pounds.', cols=[col('PlanCode', w=10), col('PlanName', w=11), col('MonthlyFee', 'money', 11), col('JoiningFee', 'money', 10), col('PeakAccess', 'bool', 11), col('ClassesPerMonth', 'int', 15), col('PoolAccess', 'bool', 10), col('MinAge', 'int', 8)],
                   rows=[list(p) for p in PLANS]))
SHEETS.append(dict(name='PriceHistory', note='Monthly fee for each plan by year. The years run across the top, so this is the sheet for HLOOKUP.', cols=[col('PlanCode', w=10)] + [col(str(y), 'money', 9) for y in PRICE_YEARS],
                   rows=[[p] + PRICE_HISTORY[p] for p in PRICE_HISTORY], header_numbers=True))
SHEETS.append(dict(name='Members', note='Everyone who joined up to the end of 2025. Status is as at 31 December 2025.', cols=[col('MemberID', w=10), col('FirstName', w=11), col('LastName', w=11), col('Gender', w=7), col('DateOfBirth', 'date', 11), col('Postcode', w=9), col('HomeSiteID', w=10), col('PlanCode', w=9), col('JoinDate', 'date', 11), col('Status', w=10), col('CancelDate', 'date', 11)],
                   rows=[[m['id'], m['first'], m['last'], m['gender'], serial(m['dob']), m['postcode'], m['site'], m['plan'], serial(m['join']), status_at(m, END25),
                          serial(m['cancel']) if (m['cancel'] and m['cancel'] <= END25) else None] for m in lesson_members]))
SHEETS.append(dict(name='Classes', note='The class timetable menu. Aqua classes only run at sites with a pool.', cols=[col('ClassID', w=8), col('ClassName', w=18), col('Category', w=12), col('DurationMins', 'int', 12), col('Intensity', w=9), col('MaxPlaces', 'int', 10)],
                   rows=[list(c[:6]) for c in CLASSES]))
def tfrac(t): hh, mm = map(int, t.split(':')); return (hh * 60 + mm) / 1440
SHEETS.append(dict(name='Bookings', note='Every class booking in 2025.', cols=[col('BookingID', w=10), col('MemberID', w=10), col('ClassID', w=8), col('SiteID', w=7), col('ClassDate', 'date', 11), col('StartTime', 'time', 9), col('Attended', 'bool', 9)],
                   rows=[[b['id'], b['member'], b['cls'], b['site'], serial(b['date']), tfrac(b['time']), b['att']] for b in bookings25]))
SHEETS.append(dict(name='Visits', note='Every gym check-in in 2025.', cols=[col('VisitID', w=9), col('MemberID', w=10), col('SiteID', w=7), col('CheckIn', 'datetime', 16), col('DurationMins', 'int', 12)],
                   rows=[[v['id'], v['member'], v['site'], round(serial(v['checkin']), 6), v['dur']] for v in visits25]))
SHEETS.append(dict(name='SiteMonths', note='One row per site per month of 2025: new members, cancellations, visits, marketing spend (£) and the average temperature (°C).', cols=[col('SiteID', w=7), col('Month', 'month', 10), col('NewMembers', 'int', 12), col('Cancellations', 'int', 12), col('Visits', 'int', 8), col('MarketingSpend', 'money0', 14), col('AvgTempC', 'num', 9, '0.0')],
                   rows=[[r[0], serial(r[1])] + r[2:] for r in sitemonths]))
SHEETS.append(dict(name='MembersRaw', note='A messy extract from the old membership system, for the Data quality and cleaning unit. Phone numbers were stored as numbers, and some rows are repeated.', cols=[col('MemberID', w=10), col('FirstName', w=12), col('LastName', w=12), col('Gender', w=8), col('DOB', w=12), col('Phone', 'int', 12, '0'), col('Postcode', w=9), col('HomeSite', w=9), col('Plan', w=9), col('MonthlyFee', 'num', 11)],
                   rows=raw_rows))

os.makedirs(os.path.join(ROOT, 'data'), exist_ok=True)
with open(os.path.join(ROOT, 'data/workbook.json'), 'w') as f:
    json.dump(dict(company='Sparta Leisure', sheets=[dict(name=s['name'], note=s['note'], cols=[{k: v for k, v in c.items() if v is not None} for c in s['cols']], rows=s['rows'], headerNumbers=s.get('header_numbers', False)) for s in SHEETS]), f, separators=(',', ':'))

# ---------------------------------------------------------------- the lesson workbook (xlsx)
NUMFMT = dict(date='dd/mm/yyyy', datetime='dd/mm/yyyy hh:mm', time='hh:mm', month='mmm yyyy', money='£#,##0.00', money0='£#,##0', int='0', num='0.00')
wb = Workbook(); ws = wb.active; ws.title = 'Read me'
F = Font(name='Arial', size=10); FB = Font(name='Arial', size=10, bold=True, color='FFFFFF'); HEAD = PatternFill('solid', fgColor='2D2A2B')
readme = [('Sparta Leisure', True), ('Practice data for the Sparta Global Data Concepts learning path. Sparta Leisure is a made-up chain of six leisure centres in the West Midlands; every person in this data is invented.', False), ('', False), ('Sheet', True)]
ws['A1'] = 'Sparta Leisure'; ws['A1'].font = Font(name='Arial', size=16, bold=True, color='D42A57')
ws['A2'] = readme[1][0]; ws['A2'].font = F
ws['A4'], ws['B4'], ws['C4'] = 'Sheet', 'Rows', 'What it holds'
for c in 'ABC': ws[c + '4'].font = FB; ws[c + '4'].fill = HEAD
for i, s in enumerate(SHEETS, 5):
    ws.cell(i, 1, s['name']).font = Font(name='Arial', size=10, bold=True)
    ws.cell(i, 2, len(s['rows'])).font = F; ws.cell(i, 3, s['note']).font = F
r = 5 + len(SHEETS) + 1
links = ['How the sheets link', 'Members.HomeSiteID, Bookings.SiteID, Visits.SiteID and SiteMonths.SiteID → Sites.SiteID',
         'Members.PlanCode and PriceHistory.PlanCode → Plans.PlanCode', 'Bookings.MemberID and Visits.MemberID → Members.MemberID', 'Bookings.ClassID → Classes.ClassID',
         'Tip: do your work on a new sheet, or in new columns to the right of the data, so you can always see the original.']
for k, t in enumerate(links):
    ws.cell(r + k, 1, t).font = Font(name='Arial', size=10, bold=(k == 0))
ws.column_dimensions['A'].width = 16; ws.column_dimensions['B'].width = 8; ws.column_dimensions['C'].width = 110
for s in SHEETS:
    sh = wb.create_sheet(s['name'])
    for j, c in enumerate(s['cols'], 1):
        hv = int(c['name']) if s.get('header_numbers') and c['name'].isdigit() else c['name']
        cell = sh.cell(1, j, hv); cell.font = FB; cell.fill = HEAD; cell.alignment = Alignment(horizontal='center')
        sh.column_dimensions[get_column_letter(j)].width = c['width'] + 2
    for i, row in enumerate(s['rows'], 2):
        for j, (v, c) in enumerate(zip(row, s['cols']), 1):
            if v is None or v == '': continue
            if c['type'] in ('date', 'month') and isinstance(v, (int, float)): v = EPOCH + timedelta(days=v)
            elif c['type'] == 'datetime': v = datetime(1899, 12, 30) + timedelta(days=v)
            elif c['type'] == 'time': v = (datetime(2000, 1, 1) + timedelta(days=v)).time()
            cell = sh.cell(i, j, v); cell.font = F
            fmt = c.get('fmt') or NUMFMT.get(c['type'])
            if fmt and not (s['name'] == 'MembersRaw' and c['type'] == 'num'): cell.number_format = fmt
    sh.freeze_panes = 'A2'
    sh.auto_filter.ref = f"A1:{get_column_letter(len(s['cols']))}{len(s['rows']) + 1}"
wb.save(os.path.join(ROOT, 'public/downloads/sparta-leisure.xlsx'))

# ---------------------------------------------------------------- the capstone export (2026, raw and messy)
S26, E26 = date(2026, 1, 1), date(2026, 9, 30)
cap_members = [m for m in members if m['join'] <= E26]
visits26 = make_visits(S26, E26, 0.255)
bookings26 = make_bookings(S26, E26, 0.62)
facts = {}

# members.csv
def fmt_dob(d, k):
    return [d.isoformat(), d.strftime('%d/%m/%Y'), d.strftime('%d %B %Y')][k]
mrows = []
for i, m in enumerate(cap_members):
    k = 0 if rng.random() < 0.8 else int(rng.integers(1, 3))
    first, last = m['first'], m['last']
    r = rng.random()
    if r < 0.06: first, last = first.upper(), last.upper()
    elif r < 0.1: first, last = first.lower(), last.lower()
    elif r < 0.14: first = ' ' + first + '  '
    gender = m['gender']
    g = rng.random()
    if g < 0.08: gender = {'F': 'Female', 'M': 'Male', 'X': 'Non-binary'}[gender]
    elif g < 0.12: gender = gender.lower()
    elif g < 0.15: gender = ''
    plan = m['plan']
    if rng.random() < 0.05: plan = {'OFF': 'Off-peak', 'STD': 'Standard', 'PREM': 'premium', 'STU': 'Student', 'SEN': 'senior'}[plan]
    site = m['site']
    if rng.random() < 0.04: site = dict((s[0], s[2]) for s in SITES)[site]
    phone = '07' + ''.join(str(int(x)) for x in rng.integers(0, 10, 9))
    if rng.random() < 0.3: phone = phone[1:]  # leading zero lost when the old system stored it as a number
    postcode = m['postcode'] if rng.random() > 0.03 else ''
    status = status_at(m, E26)
    mrows.append(dict(MemberID=m['id'], FirstName=first, LastName=last, Gender=gender, DateOfBirth=fmt_dob(m['dob'], k), Phone=phone,
                      Postcode=postcode, HomeSite=site, Plan=plan, JoinDate=m['join'].isoformat(), Status=status,
                      CancelDate=m['cancel'].isoformat() if (m['cancel'] and m['cancel'] <= E26) else ''))
bad_dob = rng.choice(len(mrows), 7, replace=False)
for j, k in enumerate(bad_dob):
    mrows[int(k)]['DateOfBirth'] = '1900-01-01' if j < 5 else '2031-05-12'
facts['membersImpossibleDob'] = 7
dups = rng.choice(len(mrows), 14, replace=False)
for k in dups: mrows.append(dict(mrows[int(k)]))
random.shuffle(mrows)
facts['membersRowsRaw'] = len(mrows); facts['membersDuplicates'] = 14; facts['membersUnique'] = len(cap_members)
facts['membersBlankGender'] = sum(1 for r in mrows if r['Gender'] == '')
facts['membersPhoneMissingZero'] = sum(1 for r in mrows if not r['Phone'].startswith('0'))
facts['membersSiteAsTown'] = sum(1 for r in mrows if not r['HomeSite'].startswith('S'))

# visits.csv: check-in and check-out times, a double export, orphans and forgotten check-outs
vrows = []
for i, v in enumerate(visits26, 1):
    out = v['checkin'] + timedelta(minutes=v['dur'])
    vrows.append(dict(VisitID=f'V26{i:05d}', MemberID=v['member'], SiteID=v['site'], CheckIn=v['checkin'], CheckOut=out))
facts['visitsTrue'] = len(vrows)
for k in rng.choice(len(vrows), 70, replace=False):  # forgot to check out: the system closes the visit at 23:59
    v = vrows[int(k)]; v['CheckOut'] = v['CheckIn'].replace(hour=23, minute=59)
facts['visitsForgotCheckout'] = 70
for k in rng.choice(len(vrows), 35, replace=False):  # clock error: check-out before check-in
    v = vrows[int(k)]; v['CheckOut'] = v['CheckIn'] - timedelta(minutes=int(rng.integers(5, 60)))
facts['visitsNegative'] = 35
for k in rng.choice(len(vrows), 22, replace=False):
    vrows[int(k)]['MemberID'] = f'M{int(rng.integers(9000, 9999))}'
facts['visitsOrphan'] = 22
dupv = [dict(vrows[int(k)]) for k in rng.choice(len(vrows), 260, replace=False)]
vrows += dupv
facts['visitsDuplicates'] = 260
for v in vrows:
    if v['SiteID'] == 'S04' and rng.random() < 0.5: v['SiteID'] = 's04'
    if v['SiteID'] == 'S06': v['SiteID'] = 'S6'
facts['visitsRowsRaw'] = len(vrows)
vrows.sort(key=lambda v: (v['CheckIn'], v['VisitID']))
def ts(d, site):  # Walsall's terminal exports UK-style text dates
    return d.strftime('%d/%m/%Y %H:%M') if site.upper() == 'S04' else d.strftime('%Y-%m-%d %H:%M')

# bookings.json
brows = []
for i, b in enumerate(bookings26, 1):
    att = b['att']
    r = rng.random()
    att_v = att if r < 0.85 else ('yes' if att else 'no') if r < 0.93 else ('Y' if att else 'N') if r < 0.97 else None
    brows.append(dict(bookingId=f'B26{i:05d}', memberId=b['member'], classId=b['cls'], siteId=b['site'],
                      classDate=b['date'].isoformat(), startTime=b['time'], attended=att_v))
facts['bookingsRows'] = len(brows); facts['bookingsAttendedBlank'] = sum(1 for b in brows if b['attended'] is None)

# monthly marketing figures (both years) so the January claim can be tested
mk = []
for (s, mo), v in sorted(spend.items(), key=lambda kv: (kv[0][1], kv[0][0])): mk.append([month_start(2025, mo).isoformat()[:7], s, v])
for (s, mo), v in sorted(spend26.items(), key=lambda kv: (kv[0][1], kv[0][0])): mk.append([month_start(2026, mo).isoformat()[:7], s, v])

def joins(y, mo): return sum(1 for m in members if m['join'].year == y and m['join'].month == mo)
facts['joinsDec2025'] = joins(2025, 12); facts['joinsJan2026'] = joins(2026, 1)
facts['joinsDec2024'] = joins(2024, 12); facts['joinsJan2025'] = joins(2025, 1)
facts['claimRise'] = round((facts['joinsJan2026'] / facts['joinsDec2025'] - 1) * 100)
facts['priorRise'] = round((facts['joinsJan2025'] / facts['joinsDec2024'] - 1) * 100)
facts['marketingJan2026'] = sum(v for (s, mo), v in spend26.items() if mo == 1)
facts['marketingJan2025'] = sum(v for (s, mo), v in spend.items() if mo == 1)

buf = io.BytesIO()
with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
    def csv_text(header, rows):
        s = io.StringIO(); w = csv.writer(s, lineterminator='\r\n'); w.writerow(header); w.writerows(rows); return s.getvalue()
    hdr = list(mrows[0].keys())
    z.writestr('sparta-leisure-capstone/members.csv', csv_text(hdr, [[r[k] for k in hdr] for r in mrows]))
    z.writestr('sparta-leisure-capstone/visits.csv', csv_text(['VisitID', 'MemberID', 'SiteID', 'CheckIn', 'CheckOut'],
               [[v['VisitID'], v['MemberID'], v['SiteID'], ts(v['CheckIn'], v['SiteID']), ts(v['CheckOut'], v['SiteID'])] for v in vrows]))
    z.writestr('sparta-leisure-capstone/bookings.json', json.dumps({'exported': '2026-10-01T06:00:00Z', 'source': 'ClassBook v3', 'bookings': brows}, indent=1))
    z.writestr('sparta-leisure-capstone/bookings.csv', csv_text(list(brows[0].keys()), [['' if v is None else v for v in b.values()] for b in brows]))
    sx = ['<?xml version="1.0" encoding="UTF-8"?>', '<sites company="Sparta Leisure">']
    for s in SITES:
        sx.append(f'  <site id="{s[0]}">\n    <name>{s[1]}</name>\n    <town>{s[2]}</town>\n    <pool>{"yes" if s[3] else "no"}</pool>\n    <capacity>{s[4]}</capacity>\n    <opened>{s[5].isoformat()}</opened>\n    <location lat="{s[6]}" lon="{s[7]}"/>\n  </site>')
    sx.append('</sites>')
    z.writestr('sparta-leisure-capstone/sites.xml', '\n'.join(sx) + '\n')
    z.writestr('sparta-leisure-capstone/plans.csv', csv_text(['PlanCode', 'PlanName', 'MonthlyFee2026', 'JoiningFee', 'PeakAccess', 'ClassesPerMonth', 'PoolAccess', 'MinAge'],
               [[p[0], p[1], round(p[2] + {'OFF': 1, 'STD': 1.5, 'PREM': 2, 'STU': 1, 'SEN': 1}[p[0]], 2), p[3], 'yes' if p[4] else 'no', p[5], 'yes' if p[6] else 'no', p[7]] for p in PLANS]))
    z.writestr('sparta-leisure-capstone/classes.csv', csv_text(['ClassID', 'ClassName', 'Category', 'DurationMins', 'Intensity', 'MaxPlaces'], [list(c[:6]) for c in CLASSES]))
    z.writestr('sparta-leisure-capstone/marketing.csv', csv_text(['Month', 'SiteID', 'MarketingSpend'], mk))
    z.writestr('sparta-leisure-capstone/README.txt', '\r\n'.join([
        'Sparta Leisure: capstone data export, 1 January to 30 September 2026',
        '',
        'members.csv     every member who joined up to 30 September 2026 (status as at 30 September 2026)',
        'visits.csv      every gym check-in and check-out',
        'bookings.json   every class booking, exported from the ClassBook system (bookings.csv is the same data as CSV)',
        'sites.xml       the six leisure centres',
        'plans.csv       membership plans and their 2026 fees',
        'classes.csv     the class menu',
        'marketing.csv   marketing spend (£) by site and month, January 2025 to September 2026',
        '',
        'This is a raw export. It has NOT been cleaned. Everyone in it is invented.', '']))
with open(os.path.join(ROOT, 'public/downloads/sparta-leisure-capstone.zip'), 'wb') as f: f.write(buf.getvalue())
with open(os.path.join(ROOT, 'data/capstone_facts.json'), 'w') as f: json.dump(facts, f, indent=1)

print('members', len(lesson_members), 'visits25', len(visits25), 'bookings25', len(bookings25), 'raw', len(raw_rows))
print('capstone', len(mrows), len(vrows), len(brows)); print(json.dumps(facts))
